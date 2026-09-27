"""Unit tests verifying the generation, integrity, and layout of the academic docx document and banner asset."""

from pathlib import Path
from docx import Document
from PIL import Image

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"
ASSETS_DIR = DOCS_DIR / "assets"
BANNER_PATH = ASSETS_DIR / "repro_radar_banner.png"
DOCX_PATH = DOCS_DIR / "repro_radar_introduction.docx"


def test_cover_banner_image_properties():
    assert BANNER_PATH.exists(), "Cover banner PNG must exist in docs/assets/"
    with Image.open(BANNER_PATH) as img:
        assert img.format == "PNG"
        assert img.width == 1200
        assert img.height == 420


def test_academic_introduction_docx_integrity():
    assert DOCX_PATH.exists(), "Academic introduction Word document (.docx) must exist in docs/"
    doc = Document(DOCX_PATH)

    # Verify paragraphs and text content
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Repro Radar" in full_text
    assert "Empirical Reproducibility" in full_text
    assert "Isolated Sandboxing" in full_text
    assert "Four-Tier" in full_text or "4-Tier" in full_text
    assert "Claerbout" in full_text

    # Verify tables (Metadata, Sandbox Specs, Rubric)
    assert len(doc.tables) >= 3, "Document should contain structured comparison and metadata tables"

    # Verify headings
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Crisis of Empirical Reproducibility" in h for h in headings)
    assert any("Container-Isolated Sandboxing" in h for h in headings)
    assert any("Four-Tier" in h for h in headings)
    assert any("References" in h for h in headings)


def _body_text(doc):
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )


def test_introduction_is_marked_draft():
    doc = Document(DOCX_PATH)
    assert doc.paragraphs[0].text.startswith("DRAFT") or any(p.text.startswith("DRAFT") for p in doc.paragraphs[:4])
    header_text = " ".join(p.text for s in doc.sections for p in s.header.paragraphs)
    assert "DRAFT" in header_text


def test_introduction_drops_unsourced_and_out_of_scope_claims():
    text = _body_text(Document(DOCX_PATH))
    for phrase in ["fewer than 35", "In 1995", "Domain-Agnostic", "domain-agnostic", "cryptography",
                   "computational economics", "Signal Processing", "resolves", "Universal Scientific"]:
        assert phrase not in text, phrase
    assert "pilot study on tabular ML benchmarks" in text


def test_every_reference_is_cited_in_the_text():
    import re

    doc = Document(DOCX_PATH)
    paragraphs = [p.text for p in doc.paragraphs]
    ref_start = next(i for i, t in enumerate(paragraphs) if "References" in t)
    body = "\n".join(paragraphs[:ref_start]) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    ref_numbers = [int(m.group(1)) for t in paragraphs[ref_start + 1:] for m in [re.match(r"\[(\d+)\]", t)] if m]
    cited = {int(n) for group in re.findall(r"\[(\d+(?:,\s*\d+)*)\]", body) for n in re.split(r",\s*", group)}
    assert ref_numbers, "references must be numbered [1], [2], ..."
    assert set(ref_numbers) == cited, (ref_numbers, sorted(cited))
